#!/usr/bin/env python3
"""Foot-check a spreadsheet: does a stated total cell equal the sum of a component range?

Footing errors (rows don't sum to the stated total) are the most common defect in generated
budget/finance sheets. Use this to (a) diagnose the exact delta before fixing, and (b) verify the
fix reconciles to zero afterward.

Reads computed cell values (openpyxl data_only). If the sheet uses formulas that openpyxl hasn't
evaluated (freshly written by openpyxl, never opened in Excel/LibreOffice), run the xlsx skill's
scripts/recalc.py first so cached values exist.

Examples:
  # Sum the country rows in E7:E28 and compare to the total cell E27
  python verify_xlsx.py budget.xlsx --sheet "Per-Country Allocation" --total E27 --range E7:E28

  # Compare against an explicit expected number instead of a total cell
  python verify_xlsx.py budget.xlsx --sheet "Dec 2025" --range B3:B12 --expect 5600

  # Just print the sum of a range
  python verify_xlsx.py budget.xlsx --sheet Sheet1 --range B3:B12
"""
import argparse
import os
import re
import sys
import zipfile

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries, get_column_letter
except ImportError:
    print("ERROR: openpyxl required (pip install openpyxl)", file=sys.stderr)
    sys.exit(2)


def numeric(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace("$", "").replace(",", "").replace("(", "-").replace(")", "")
        try:
            return float(s)
        except ValueError:
            return None
    return None



# --------------------------------------------------------------------------------------------------
# Formula cells with no cached value.
#
# An xlsx stores each formula twice: the formula itself, and the last computed result Excel cached
# next to it. Anything that reads VALUES, including AutoQC, reads the cache. It does not calculate.
#
# openpyxl does not calculate either. When it writes a file it emits `<f>SUM(A1:A2)</f><v></v>`, an
# explicitly EMPTY cached value. So the formula survives, the number does not, and a grader reading
# that sheet sees blanks where the totals should be. Nothing errors, the file opens fine in Excel
# because Excel recalculates on open, and the defect is invisible until AutoQC reports empty cells.
#
# This finds it. Detection is the missing half: the fixes are easy once you know.
# --------------------------------------------------------------------------------------------------
CELL = re.compile(rb"<c\b[^>]*r=\"([A-Z]+\d+)\"[^>]*>(.*?)</c>", re.S)
HAS_F = re.compile(rb"<f[ >]")
HAS_V = re.compile(rb"<v>\s*[^<\s]", re.S)


def stale_formula_cells(path):
    """[(sheet_xml_name, cell_ref)] for formula cells whose cached value is empty or absent."""
    out = []
    try:
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                    continue
                blob = z.read(name)
                for m in CELL.finditer(blob):
                    ref, body = m.group(1).decode(), m.group(2)
                    if HAS_F.search(body) and not HAS_V.search(body):
                        out.append((name.split("/")[-1], ref))
    except Exception:  # noqa: BLE001
        return []
    return out


def cmd_scan(root):
    files = []
    if os.path.isdir(root):
        for dp, _, fns in os.walk(root):
            for fn in fns:
                if fn.lower().endswith((".xlsx", ".xlsm")) and not fn.startswith("~$"):
                    p = os.path.join(dp, fn)
                    if "_qc_backup" not in p.replace("\\", "/"):
                        files.append(p)
    else:
        files = [root]

    bad = []
    for p in sorted(files):
        cells = stale_formula_cells(p)
        if cells:
            bad.append((p, cells))

    print("=" * 78)
    print("FORMULA CELLS WITH NO CACHED VALUE")
    print("=" * 78)
    print(f"\nChecked {len(files)} workbook(s).\n")
    if not bad:
        print("  None. Every formula carries its computed result, which is what a grader reads.")
        return 0

    total = sum(len(c) for _, c in bad)
    print(f"  {total} formula cell(s) across {len(bad)} workbook(s) have a formula but no number.\n")
    for p, cells in bad[:15]:
        rel = os.path.relpath(p, root if os.path.isdir(root) else os.path.dirname(root))
        shown = ", ".join(f"{s}!{r}" for s, r in cells[:8])
        more = f"  (+{len(cells) - 8} more)" if len(cells) > 8 else ""
        print(f"    {rel}   {len(cells)} cell(s)")
        print(f"      {shown}{more}")
    if len(bad) > 15:
        print(f"    ... and {len(bad) - 15} more workbook(s)")

    print("""
  WHY THIS MATTERS
  AutoQC reads cached values, it does not calculate. These cells will read as empty to the grader
  even though the workbook opens correctly in Excel, because Excel recalculates on open and the
  grader does not. This is usually self-inflicted: openpyxl writes an empty <v> whenever it saves.

  THREE WAYS TO FIX, easiest first:
  1. Open the workbook in Excel, press Ctrl+S, close it. Excel recalculates on open and writes the
     cached values back. No install needed and most builders already have Excel.
  2. If you have LibreOffice:
       soffice --headless --convert-to xlsx --outdir <same folder> <file.xlsx>
     It opens, recalculates and re-saves. Only worth installing if you have no Excel at all.
  3. Patch the cached value in the sheet XML alongside the formula, so no recalculation is needed.
     Slower and fiddlier, but it is the only option with neither Excel nor LibreOffice.

  BEST OF ALL: do not create it. When editing a workbook, write the new number into the cached
  value as well as the formula, and this never appears.""")
    return len(bad)


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    # --scan is a whole separate mode: it takes a folder and no other arguments. Handled before
    # the parser runs so the foot-check's required arguments do not fire on a scan.
    if "--scan" in sys.argv:
        i = sys.argv.index("--scan")
        if i + 1 >= len(sys.argv):
            print("ERROR: --scan needs a folder, e.g. --scan .", file=sys.stderr)
            sys.exit(2)
        sys.exit(1 if cmd_scan(sys.argv[i + 1]) else 0)

    p.add_argument("file")
    p.add_argument("--scan", metavar="DIR",
                   help="separate mode: sweep a folder for formula cells that lost their cached "
                        "value (the openpyxl round-trip defect). Takes no other arguments.")
    p.add_argument("--sheet", required=True)
    p.add_argument("--range", required=True, dest="rng",
                   help="component range to sum, e.g. E7:E28")
    p.add_argument("--total", default=None,
                   help="cell holding the stated total, e.g. E27 (excluded from the sum "
                        "automatically if inside --range)")
    p.add_argument("--expect", type=float, default=None,
                   help="expected total as a literal number (alternative to --total)")
    p.add_argument("--tol", type=float, default=0.01,
                   help="tolerance for the reconciliation (default 0.01)")
    args = p.parse_args()

    try:
        wb = load_workbook(args.file, data_only=True)
    except Exception as e:
        print(f"ERROR: could not open {args.file!r}: {e}", file=sys.stderr)
        sys.exit(2)
    if args.sheet not in wb.sheetnames:
        print(f"ERROR: sheet {args.sheet!r} not found. Sheets: {wb.sheetnames}", file=sys.stderr)
        sys.exit(2)
    ws = wb[args.sheet]

    try:
        min_c, min_r, max_c, max_r = range_boundaries(args.rng)
    except ValueError as e:
        print(f"ERROR: invalid --range {args.rng!r}: {e}", file=sys.stderr)
        sys.exit(2)
    if None in (min_c, min_r, max_c, max_r):
        print("ERROR: --range must be a bounded rectangle like B2:B28, not a whole column/row",
              file=sys.stderr)
        sys.exit(2)
    if min_r > max_r or min_c > max_c:
        print(f"ERROR: --range {args.rng!r} is reversed (end before start)", file=sys.stderr)
        sys.exit(2)
    total_coord = args.total.upper() if args.total else None
    if total_coord:
        try:
            range_boundaries(total_coord)
        except ValueError as e:
            print(f"ERROR: invalid --total {args.total!r}: {e}", file=sys.stderr)
            sys.exit(2)

    components, skipped, total_in_range = [], [], False
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            coord = f"{get_column_letter(c)}{r}"
            if total_coord and coord == total_coord:
                total_in_range = True
                continue
            v = ws[coord].value
            n = numeric(v)
            if n is None:
                if v is not None:
                    skipped.append((coord, v))
            else:
                components.append((coord, n))

    csum = sum(n for _, n in components)
    print(f"Sheet:            {args.sheet}")
    print(f"Component range:  {args.rng}  ({len(components)} numeric cells summed"
          + (f", total cell {total_coord} excluded" if total_in_range else "") + ")")
    if skipped:
        preview = ", ".join(f"{c}={v!r}" for c, v in skipped[:5])
        print(f"Non-numeric cells skipped ({len(skipped)}): {preview}"
              + (" ..." if len(skipped) > 5 else ""))
    print(f"Sum of components: {csum:,.2f}")

    stated = None
    total_requested = total_coord is not None or args.expect is not None
    if total_coord:
        stated = numeric(ws[total_coord].value)
        print(f"Stated total ({total_coord}): "
              + (f"{stated:,.2f}" if stated is not None else f"{ws[total_coord].value!r} (non-numeric)"))
    if args.expect is not None:
        stated = args.expect
        print(f"Expected total:    {stated:,.2f}")

    if stated is None:
        if total_requested:
            print("\nERROR: stated total is non-numeric (formula cache empty? run recalc.py first)",
                  file=sys.stderr)
            sys.exit(2)
        print("\n(no --total/--expect given; sum only)")
        return
    delta = csum - stated
    if abs(delta) <= args.tol:
        print(f"\nRECONCILES ✓  delta = {delta:,.2f}")
        sys.exit(0)
    else:
        print(f"\nMISMATCH ✗   delta = {delta:,.2f}   (components {'exceed' if delta>0 else 'fall short of'} stated by {abs(delta):,.2f})")
        sys.exit(1)


if __name__ == "__main__":
    main()
